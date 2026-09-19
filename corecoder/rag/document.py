from __future__ import annotations

import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path
from typing import ClassVar


@dataclass
class Document:
    """原始文档。"""

    content: str
    source: str = ""
    metadata: dict = field(default_factory=dict)


@dataclass
class DocumentChunk:
    """文档切分后的 Chunk。"""

    content: str
    source: str
    metadata: dict = field(default_factory=dict)


class BaseChunker(ABC):
    """文档切分器基类。"""

    def __init__(
        self,
        chunk_size: int = 800,
        chunk_overlap: int = 100,
    ):
        if chunk_size <= 0:
            raise ValueError("chunk_size 必须大于 0")

        if chunk_overlap < 0:
            raise ValueError("chunk_overlap 不能小于 0")

        if chunk_overlap >= chunk_size:
            raise ValueError(
                "chunk_overlap 必须小于 chunk_size"
            )

        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    @abstractmethod
    def split(self, document: Document) -> list[DocumentChunk]:
        """切分文档。"""
        pass

    def _create_chunk(
        self,
        document: Document,
        content: str,
        chunk_index: int,
        extra_metadata: dict | None = None,
    ) -> DocumentChunk:

        metadata = dict(document.metadata)
        metadata["chunk_index"] = chunk_index

        if extra_metadata:
            metadata.update(extra_metadata)

        return DocumentChunk(
            content=content.strip(),
            source=document.source,
            metadata=metadata,
        )


class TextChunker(BaseChunker):
    """普通文本切分器。"""

    def split(
        self,
        document: Document,
    ) -> list[DocumentChunk]:

        parts = self._recursive_split(
            document.content,
            ["\n\n", "\n", " ", ""],
        )

        return [
            self._create_chunk(
                document,
                content,
                index,
            )
            for index, content in enumerate(parts)
            if content.strip()
        ]

    def _recursive_split(
        self,
        text: str,
        separators: list[str],
    ) -> list[str]:

        text = text.strip()

        if not text:
            return []

        if len(text) <= self.chunk_size:
            return [text]

        separator = separators[0]

        if separator == "":
            return self._hard_split(text)

        parts = [
            part.strip()
            for part in text.split(separator)
            if part.strip()
        ]

        chunks = []
        current = ""

        for part in parts:
            candidate = (
                part
                if not current
                else current + separator + part
            )

            if len(candidate) <= self.chunk_size:
                current = candidate
                continue

            if current:
                chunks.append(current)
                current = ""

            if len(part) > self.chunk_size:
                if len(separators) > 1:
                    chunks.extend(
                        self._recursive_split(
                            part,
                            separators[1:],
                        )
                    )
                else:
                    chunks.extend(
                        self._hard_split(part)
                    )
            else:
                current = part

        if current:
            chunks.append(current)

        return self._add_overlap(chunks)

    def _add_overlap(
        self,
        chunks: list[str],
    ) -> list[str]:

        if self.chunk_overlap == 0:
            return chunks

        result = [chunks[0]]

        for chunk in chunks[1:]:
            previous = result[-1]
            overlap = previous[-self.chunk_overlap:]

            combined = overlap + "\n" + chunk

            if len(combined) <= self.chunk_size:
                result.append(combined)
            else:
                result.append(chunk)

        return result

    def _hard_split(
        self,
        text: str,
    ) -> list[str]:

        chunks = []
        start = 0

        while start < len(text):
            end = start + self.chunk_size

            chunks.append(
                text[start:end].strip()
            )

            if end >= len(text):
                break

            start = end - self.chunk_overlap

        return chunks


class MarkdownChunker(TextChunker):
    """Markdown 文档切分器。"""

    HEADING_PATTERN = re.compile(
        r"^(#{1,6})\s+(.+?)\s*$"
    )

    def split(
        self,
        document: Document,
    ) -> list[DocumentChunk]:

        sections = self._split_sections(
            document.content
        )

        chunks = []

        for section in sections:
            section_chunks = self._recursive_split(
                section["content"],
                ["\n\n", "\n", " ", ""],
            )

            for content in section_chunks:
                if not content.strip():
                    continue

                chunks.append(
                    self._create_chunk(
                        document,
                        content,
                        len(chunks),
                        {
                            "section": section["title"],
                            "section_level": section["level"],
                        },
                    )
                )

        return chunks

    def _split_sections(
        self,
        text: str,
    ) -> list[dict]:

        lines = text.splitlines()

        sections = []
        current_title = "ROOT"
        current_level = 0
        current_lines = []

        for line in lines:
            match = self.HEADING_PATTERN.match(line)

            if match:
                if current_lines:
                    content = "\n".join(
                        current_lines
                    ).strip()

                    if content:
                        sections.append(
                            {
                                "title": current_title,
                                "level": current_level,
                                "content": content,
                            }
                        )

                current_level = len(match.group(1))
                current_title = match.group(2)
                current_lines = []
            else:
                current_lines.append(line)

        if current_lines:
            content = "\n".join(
                current_lines
            ).strip()

            if content:
                sections.append(
                    {
                        "title": current_title,
                        "level": current_level,
                        "content": content,
                    }
                )

        return sections


class CodeChunker(BaseChunker):
    """代码文件切分器。"""

    CODE_BLOCK_PATTERN = re.compile(
        r"^(class\s+\w+|"
        r"def\s+\w+|"
        r"async\s+def\s+\w+|"
        r"function\s+\w+|"
        r"public\s+(?:static\s+)?[\w<>\[\]]+\s+\w+\s*\(|"
        r"private\s+(?:static\s+)?[\w<>\[\]]+\s+\w+\s*\(|"
        r"protected\s+(?:static\s+)?[\w<>\[\]]+\s+\w+\s*\()"
    )

    def split(
        self,
        document: Document,
    ) -> list[DocumentChunk]:

        blocks = self._split_code_blocks(
            document.content
        )

        chunks = []

        for block in blocks:
            block_chunks = self._split_block(block)

            for content in block_chunks:
                if not content.strip():
                    continue

                chunks.append(
                    self._create_chunk(
                        document,
                        content,
                        len(chunks),
                        {
                            "chunk_type": block["type"],
                            "symbol": block["symbol"],
                        },
                    )
                )

        return chunks

    def _split_code_blocks(
        self,
        text: str,
    ) -> list[dict]:

        lines = text.splitlines()

        blocks = []
        current_lines = []
        current_symbol = "module"
        current_type = "module"

        for line in lines:
            match = self.CODE_BLOCK_PATTERN.match(
                line.strip()
            )

            if match and current_lines:
                blocks.append(
                    {
                        "type": current_type,
                        "symbol": current_symbol,
                        "content": "\n".join(
                            current_lines
                        ).strip(),
                    }
                )

                current_lines = []

            if match:
                current_symbol = match.group(0)
                current_type = self._detect_code_type(
                    line
                )

            current_lines.append(line)

        if current_lines:
            blocks.append(
                {
                    "type": current_type,
                    "symbol": current_symbol,
                    "content": "\n".join(
                        current_lines
                    ).strip(),
                }
            )

        return [
            block
            for block in blocks
            if block["content"]
        ]

    def _detect_code_type(
        self,
        line: str,
    ) -> str:

        line = line.strip()

        if line.startswith("class "):
            return "class"

        if (
            line.startswith("def ")
            or line.startswith("async def ")
        ):
            return "function"

        if line.startswith("function "):
            return "function"

        return "method"

    def _split_block(
        self,
        block: dict,
    ) -> list[str]:

        text = block["content"]

        if len(text) <= self.chunk_size:
            return [text]

        lines = text.splitlines()

        chunks = []
        current = ""

        for line in lines:
            candidate = (
                line
                if not current
                else current + "\n" + line
            )

            if len(candidate) <= self.chunk_size:
                current = candidate
                continue

            if current:
                chunks.append(current)
                current = ""

            if len(line) > self.chunk_size:
                chunks.extend(
                    self._hard_split(line)
                )
            else:
                current = line

        if current:
            chunks.append(current)

        return self._add_code_overlap(chunks)

    def _add_code_overlap(
        self,
        chunks: list[str],
    ) -> list[str]:

        if (
            self.chunk_overlap == 0
            or len(chunks) <= 1
        ):
            return chunks

        result = [chunks[0]]

        for chunk in chunks[1:]:
            previous_lines = result[-1].splitlines()

            overlap_lines = []
            length = 0

            for line in reversed(previous_lines):
                line_length = len(line) + 1

                if length + line_length > self.chunk_overlap:
                    break

                overlap_lines.insert(0, line)
                length += line_length

            overlap = "\n".join(overlap_lines)

            if overlap:
                combined = overlap + "\n" + chunk

                if len(combined) <= self.chunk_size:
                    result.append(combined)
                else:
                    result.append(chunk)
            else:
                result.append(chunk)

        return result

    def _hard_split(
        self,
        text: str,
    ) -> list[str]:

        chunks = []
        start = 0

        while start < len(text):
            end = start + self.chunk_size

            chunks.append(
                text[start:end].strip()
            )

            if end >= len(text):
                break

            start = end - self.chunk_overlap

        return chunks


class ChunkerFactory:
    """根据文件类型创建 Chunker。"""

    MARKDOWN_SUFFIXES: ClassVar[
        set[str]
    ] = {
        ".md",
        ".markdown",
        ".docx",
    }

    TEXT_SUFFIXES: ClassVar[
        set[str]
    ] = {
        ".txt",
        ".text",
        ".pdf",
    }

    CODE_SUFFIXES: ClassVar[
        set[str]
    ] = {
        ".py",
        ".java",
        ".js",
        ".ts",
        ".go",
        ".cpp",
        ".c",
        ".h",
        ".hpp",
        ".rs",
    }

    @classmethod
    def create(
        cls,
        file_path: str,
        chunk_size: int = 800,
        chunk_overlap: int = 100,
    ) -> BaseChunker:
        suffix = (
            Path(file_path)
            .suffix
            .lower()
        )

        kwargs = {
            "chunk_size": chunk_size,
            "chunk_overlap": (
                chunk_overlap
            ),
        }

        if (
            suffix
            in cls.MARKDOWN_SUFFIXES
        ):
            return MarkdownChunker(
                **kwargs
            )

        if (
            suffix
            in cls.CODE_SUFFIXES
        ):
            return CodeChunker(
                **kwargs
            )

        return TextChunker(
            **kwargs
        )

    @classmethod
    def create(
        cls,
        file_path: str,
        chunk_size: int = 800,
        chunk_overlap: int = 100,
    ) -> BaseChunker:

        suffix = Path(file_path).suffix.lower()

        kwargs = {
            "chunk_size": chunk_size,
            "chunk_overlap": chunk_overlap,
        }

        if suffix in cls.MARKDOWN_SUFFIXES:
            return MarkdownChunker(**kwargs)

        if suffix in cls.CODE_SUFFIXES:
            return CodeChunker(**kwargs)

        return TextChunker(**kwargs)


def load_document(
    file_path: str,
) -> Document:
    """根据文件类型加载文档。"""

    path = Path(
        file_path
    )

    if not path.exists():
        raise FileNotFoundError(
            f"文件不存在: {file_path}"
        )

    if not path.is_file():
        raise ValueError(
            f"不是文件: {file_path}"
        )

    suffix = (
        path.suffix.lower()
    )

    supported_suffixes = (
        ChunkerFactory.MARKDOWN_SUFFIXES
        | ChunkerFactory.TEXT_SUFFIXES
        | ChunkerFactory.CODE_SUFFIXES
    )

    if (
        suffix
        not in supported_suffixes
    ):
        raise ValueError(
            "不支持的文件类型: "
            f"{suffix or '无扩展名'}"
        )

    extra_metadata = {}

    if suffix == ".pdf":
        (
            content,
            extra_metadata,
        ) = _load_pdf(path)

    elif suffix == ".docx":
        (
            content,
            extra_metadata,
        ) = _load_docx(path)

    else:
        content = path.read_text(
            encoding="utf-8"
        )

    if not content.strip():
        raise ValueError(
            "文件中没有可提取的文本: "
            f"{path.name}"
        )

    return Document(
        content=content,
        source=str(path),
        metadata={
            "file_name": path.name,
            "file_type": suffix,
            **extra_metadata,
        },
    )

def _load_pdf(
    path: Path,
) -> tuple[str, dict]:
    """提取文字型 PDF，并保留页码边界。"""

    from pypdf import PdfReader

    reader = PdfReader(
        str(path)
    )

    pages = []

    for (
        page_number,
        page,
    ) in enumerate(
        reader.pages,
        1,
    ):
        text = (
            page.extract_text()
            or ""
        ).strip()

        if text:
            pages.append(
                f"[第 {page_number} 页]\n"
                f"{text}"
            )

    return (
        "\n\n".join(pages),
        {
            "page_count": (
                len(reader.pages)
            ),
        },
    )

def _load_docx(
    path: Path,
) -> tuple[str, dict]:
    """提取 DOCX 标题、段落和表格。"""

    from docx import (
        Document as DocxDocument,
    )

    document = DocxDocument(
        str(path)
    )

    blocks = []

    for paragraph in (
        document.paragraphs
    ):
        text = (
            paragraph.text.strip()
        )

        if not text:
            continue

        style_name = (
            paragraph.style.name
            if paragraph.style
            else ""
        )

        if style_name.startswith(
            "Heading "
        ):
            try:
                level = int(
                    style_name.removeprefix(
                        "Heading "
                    )
                )
            except ValueError:
                level = 1

            level = min(
                max(level, 1),
                6,
            )

            blocks.append(
                f"{'#' * level} "
                f"{text}"
            )
        else:
            blocks.append(
                text
            )

    for (
        table_index,
        table,
    ) in enumerate(
        document.tables,
        1,
    ):
        rows = []

        for row in table.rows:
            values = [
                cell.text.strip()
                for cell
                in row.cells
            ]

            rows.append(
                " | ".join(values)
            )

        if rows:
            blocks.append(
                f"## 表格 {table_index}\n"
                + "\n".join(rows)
            )

    return (
        "\n\n".join(blocks),
        {
            "paragraph_count": (
                len(
                    document.paragraphs
                )
            ),
            "table_count": (
                len(
                    document.tables
                )
            ),
        },
    )

def split_document(
    document: Document,
    chunk_size: int = 800,
    chunk_overlap: int = 100,
) -> list[DocumentChunk]:
    """根据文件类型自动选择切分器。"""

    chunker = ChunkerFactory.create(
        document.source,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )

    return chunker.split(document)

